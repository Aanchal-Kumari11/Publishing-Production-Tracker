#!/usr/bin/env python3
"""
Publishing Production Tracker

Reads a CSV containing book-production tasks and creates:
  1. publishing_tasks_report.xlsx - formatted Excel workbook with color-coded status/reminders
  2. publishing_summary.txt - a concise production summary

The formatted Excel workbook requires openpyxl: pip install openpyxl.
A text summary is also generated. If openpyxl is unavailable, the Excel workbook will not be created.

Example:
    python publishing_production_tracker.py input_tasks.csv
    python publishing_production_tracker.py input_tasks.csv --today 2026-10-14
    python publishing_production_tracker.py input_tasks.csv --output-dir reports

Suggested CSV columns:
    book_title,task,deadline,status,priority,assigned_to,stage,dependency,notes

Column names are flexible; common alternatives such as "book", "task_name",
"due_date", "owner", and "progress" are recognized. Existing rows are never
removed just because a book title repeats. Exact duplicate rows are retained too.
"""

from __future__ import annotations

import argparse
import csv
from collections import Counter, defaultdict
from datetime import date, datetime
from pathlib import Path
import re
import sys


COMPLETED_STATUSES = {
    "approved", "complete", "completed", "closed", "done",
    "published", "delivered", "cancelled", "canceled",
}

STATUS_ALIASES = {
    "in progress": "In Progress",
    "in-progress": "In Progress",
    "inprogress": "In Progress",
    "not started": "Not Started",
    "not-started": "Not Started",
    "todo": "Not Started",
    "to do": "Not Started",
    "in review": "In Review",
    "in-review": "In Review",
    "review": "In Review",
    "pending": "Pending",
    "blocked": "Blocked",
    "on hold": "On Hold",
    "on-hold": "On Hold",
    "done": "Completed",
    "complete": "Completed",
    "completed": "Completed",
    "approved": "Approved",
    "published": "Published",
    "cancelled": "Cancelled",
    "canceled": "Cancelled",
}

PRIORITY_RANK = {
    "critical": 0,
    "urgent": 0,
    "high": 1,
    "medium": 2,
    "normal": 2,
    "low": 3,
    "": 2,
}

COLUMN_ALIASES = {
    "book_title": ("book_title", "book", "book_name", "title", "project", "project_name"),
    "task": ("task", "task_name", "activity", "deliverable", "work_item", "item", "asset_name"),
    "deadline": ("deadline", "due_date", "due", "target_date", "completion_date", "publication_deadline"),
    "status": ("status", "task_status", "progress", "state"),
    "priority": ("priority", "urgency", "importance"),
    "assigned_to": ("assigned_to", "owner", "assignee", "responsible", "person"),
    "stage": ("stage", "production_stage", "department", "work_stage"),
    "dependency": ("dependency", "depends_on", "blocked_by", "prerequisite"),
    "notes": ("notes", "comments", "blocker", "remarks", "description"),
    "task_id": ("task_id", "id", "work_id", "asset_id"),
}


def normalize_header(value: str) -> str:
    value = (value or "").strip().lower()
    value = re.sub(r"[^a-z0-9]+", "_", value)
    return value.strip("_")


def clean_text(value: object) -> str:
    return " ".join(str(value or "").strip().split())


def parse_date(value: str) -> date | None:
    value = clean_text(value)
    if not value:
        return None

    # Common spreadsheet and CSV date formats.
    formats = (
        "%Y-%m-%d",
        "%d-%m-%Y",
        "%d/%m/%Y",
        "%m/%d/%Y",
        "%d %b %Y",
        "%d %B %Y",
        "%Y/%m/%d",
        "%d-%b-%Y",
    )
    for fmt in formats:
        try:
            return datetime.strptime(value, fmt).date()
        except ValueError:
            continue
    return None


def canonical_status(value: str) -> str:
    raw = clean_text(value)
    key = raw.lower()
    return STATUS_ALIASES.get(key, raw.title() if raw else "Not Specified")


def is_completed(status: str) -> bool:
    return clean_text(status).lower() in COMPLETED_STATUSES


def canonical_priority(value: str) -> str:
    raw = clean_text(value).lower()
    if raw in {"urgent", "critical"}:
        return "Critical"
    if raw == "high":
        return "High"
    if raw in {"medium", "normal", ""}:
        return "Medium" if raw else "Unspecified"
    if raw == "low":
        return "Low"
    return raw.title()


def find_columns(headers: list[str]) -> dict[str, str | None]:
    normalized_to_original = {normalize_header(h): h for h in headers}
    found: dict[str, str | None] = {}
    for field, aliases in COLUMN_ALIASES.items():
        found[field] = next(
            (normalized_to_original[a] for a in aliases if a in normalized_to_original),
            None,
        )
    return found


def reminder_label(days_left: int | None, completed: bool) -> str:
    if completed:
        return "No reminder - completed"
    if days_left is None:
        return "Check deadline - missing/invalid date"
    if days_left < 0:
        overdue_by = abs(days_left)
        return f"OVERDUE by {overdue_by} day(s)"
    if days_left == 0:
        return "DUE TODAY"
    if days_left == 1:
        return "Due tomorrow"
    if days_left == 2:
        return "Due in 2 days"
    if days_left == 3:
        return "Due in 3 days"
    return f"Due in {days_left} days"


def urgency_rank(days_left: int | None, completed: bool) -> int:
    if completed:
        return 9
    if days_left is None:
        return 7
    if days_left < 0:
        return 0
    if days_left == 0:
        return 1
    if days_left <= 3:
        return 2
    if days_left <= 7:
        return 3
    return 4


def read_tasks(csv_path: Path, today: date) -> tuple[list[dict[str, str]], list[str]]:
    if not csv_path.exists():
        raise FileNotFoundError(f"Input CSV not found: {csv_path}")

    with csv_path.open("r", newline="", encoding="utf-8-sig") as handle:
        reader = csv.DictReader(handle)
        if not reader.fieldnames:
            raise ValueError("The input CSV is empty or has no header row.")

        columns = find_columns(reader.fieldnames)
        if not columns["book_title"] or not columns["task"]:
            raise ValueError(
                "Could not identify required columns for book title and task. "
                "Use headers such as 'book_title' and 'task' (or 'book' and 'task_name')."
            )

        tasks: list[dict[str, str]] = []
        warnings: list[str] = []
        for line_number, raw in enumerate(reader, start=2):
            if not any(clean_text(v) for v in raw.values()):
                # Fully blank lines can be ignored; no task information is lost.
                continue

            row = {normalize_header(k): clean_text(v) for k, v in raw.items() if k is not None}
            def get(field: str, default: str = "") -> str:
                original_header = columns.get(field)
                return clean_text(raw.get(original_header, "")) if original_header else default

            book = get("book_title") or "Book title missing"
            task_name = get("task") or "Task description missing"
            raw_deadline = get("deadline")
            parsed_deadline = parse_date(raw_deadline)
            status = canonical_status(get("status"))
            priority = canonical_priority(get("priority"))
            owner = get("assigned_to") or "Unassigned"
            stage = get("stage") or "Unspecified"
            dependency = get("dependency")
            notes = get("notes")
            task_id = get("task_id")

            if raw_deadline and parsed_deadline is None:
                warnings.append(
                    f"Row {line_number}: could not parse deadline '{raw_deadline}' for "
                    f"'{book} - {task_name}'. It will be flagged for review."
                )

            completed = is_completed(status)
            days_left = (parsed_deadline - today).days if parsed_deadline else None
            reminder = reminder_label(days_left, completed)

            if completed:
                priority_order = 9
            else:
                priority_order = PRIORITY_RANK.get(priority.lower(), 2)

            tasks.append({
                "book_title": book,
                "task_id": task_id,
                "task": task_name,
                "stage": stage,
                "assigned_to": owner,
                "deadline": parsed_deadline.isoformat() if parsed_deadline else raw_deadline,
                "days_until_deadline": str(days_left) if days_left is not None else "",
                "reminder": reminder,
                "priority": priority,
                "status": status,
                "dependency": dependency,
                "notes": notes,
                "_completed": "yes" if completed else "no",
                "_urgency_rank": str(urgency_rank(days_left, completed)),
                "_priority_rank": str(priority_order),
                "_book_deadline": parsed_deadline.isoformat() if parsed_deadline else "9999-12-31",
                "_original_row": str(line_number),
            })

    return tasks, warnings


def sort_tasks(tasks: list[dict[str, str]]) -> list[dict[str, str]]:
    # Group books together, ordering books by their most urgent active task.
    # Within each book, tasks are ordered by urgency, priority, then deadline.
    book_keys: dict[str, tuple[int, int, str]] = {}
    for task in tasks:
        if task["_completed"] == "yes":
            continue
        days = int(task["days_until_deadline"]) if task["days_until_deadline"] else None
        candidate = (
            urgency_rank(days, False),
            PRIORITY_RANK.get(task["priority"].lower(), 2),
            task["_book_deadline"],
        )
        book = task["book_title"].casefold()
        if book not in book_keys or candidate < book_keys[book]:
            book_keys[book] = candidate

    def key(task: dict[str, str]):
        book = task["book_title"].casefold()
        book_rank = book_keys.get(book, (9, 9, "9999-12-31"))
        days = int(task["days_until_deadline"]) if task["days_until_deadline"] else None
        return (
            book_rank[0],
            book_rank[1],
            book_rank[2],
            book,
            urgency_rank(days, task["_completed"] == "yes"),
            PRIORITY_RANK.get(task["priority"].lower(), 2),
            task["_book_deadline"],
            task["_original_row"],
        )

    return sorted(tasks, key=key)


def write_summary(tasks: list[dict[str, str]], warnings: list[str], output_path: Path, today: date) -> str:
    active = [t for t in tasks if t["_completed"] != "yes"]
    completed = [t for t in tasks if t["_completed"] == "yes"]
    overdue = [t for t in active if t["days_until_deadline"] and int(t["days_until_deadline"]) < 0]
    due_today = [t for t in active if t["days_until_deadline"] == "0"]
    due_soon = [t for t in active if t["days_until_deadline"] in {"1", "2", "3"}]
    missing_dates = [t for t in active if not t["days_until_deadline"]]
    blocked = [t for t in active if t["status"].lower() in {"blocked", "on hold"} or t["dependency"]]

    books: dict[str, list[dict[str, str]]] = defaultdict(list)
    for task in tasks:
        books[task["book_title"]].append(task)

    lines = [
        "BOOK PUBLISHING PRODUCTION SUMMARY",
        "=" * 42,
        f"Report date: {today.isoformat()}",
        f"Tasks retained (including repeated book names and duplicate rows): {len(tasks)}",
        f"Books represented: {len(books)}",
        f"Active tasks: {len(active)}",
        f"Completed/cancelled tasks: {len(completed)}",
        f"Overdue tasks: {len(overdue)}",
        f"Due today: {len(due_today)}",
        f"Due within 3 days: {len(due_soon)}",
        f"Missing/invalid deadlines: {len(missing_dates)}",
        f"Blocked/on-hold or dependency noted: {len(blocked)}",
        "",
        "PRIORITY FOLLOW-UP",
        "-" * 42,
    ]

    follow_up = overdue + due_today + due_soon + missing_dates + [
        t for t in blocked if t not in overdue and t not in due_today and t not in due_soon and t not in missing_dates
    ]
    seen_rows = set()
    for task in follow_up:
        unique_row = task["_original_row"]
        if unique_row in seen_rows:
            continue
        seen_rows.add(unique_row)
        lines.append(
            f"- {task['book_title']} | {task['task']} | {task['deadline'] or 'No valid deadline'} "
            f"| {task['reminder']} | {task['status']} | Owner: {task['assigned_to']}"
        )
    if not follow_up:
        lines.append("- No immediate follow-ups flagged.")

    lines.extend(["", "TASKS BY BOOK", "-" * 42])
    for book, book_tasks in sorted(books.items(), key=lambda item: item[0].casefold()):
        open_count = sum(t["_completed"] != "yes" for t in book_tasks)
        done_count = len(book_tasks) - open_count
        lines.append(f"- {book}: {len(book_tasks)} task(s), {open_count} active, {done_count} completed")

    if warnings:
        lines.extend(["", "DATA QUALITY WARNINGS", "-" * 42])
        lines.extend(f"- {warning}" for warning in warnings)
    else:
        lines.extend(["", "DATA QUALITY", "-" * 42, "- No unparseable dates found."])

    output_path.write_text("\n".join(lines) + "\n", encoding="utf-8")
    return "\n".join(lines)



def write_excel_report(tasks: list[dict[str, str]], warnings: list[str], output_path: Path, today: date) -> bool:
    """Write a styled Excel workbook. Returns False if openpyxl is not installed."""
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill, Border, Side, Alignment
        from openpyxl.utils import get_column_letter
    except ImportError:
        return False

    navy = "203864"
    blue = "4472C4"
    white = "FFFFFF"
    light_blue = "D9EAF7"
    light_grey = "F2F4F7"
    red_fill = "F4CCCC"
    red_font = "990000"
    orange_fill = "FCE4D6"
    orange_font = "C65911"
    yellow_fill = "FFF2CC"
    yellow_font = "7F6000"
    green_fill = "D9EAD3"
    green_font = "38761D"
    purple_fill = "E4DFEC"
    purple_font = "674EA7"
    border_color = "D9E1F2"

    header_fill = PatternFill("solid", fgColor=navy)
    header_font = Font(color=white, bold=True)
    thin_border = Border(bottom=Side(style="thin", color=border_color))

    wb = Workbook()
    ws = wb.active
    ws.title = "Task Tracker"

    columns = [
        ("Book Title", "book_title"),
        ("Task ID", "task_id"),
        ("Task / Deliverable", "task"),
        ("Production Stage", "stage"),
        ("Assigned To", "assigned_to"),
        ("Deadline", "deadline"),
        ("Days Left", "days_until_deadline"),
        ("Reminder", "reminder"),
        ("Priority", "priority"),
        ("Status", "status"),
        ("Dependency", "dependency"),
        ("Notes", "notes"),
    ]

    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(columns))
    title_cell = ws.cell(1, 1, "BOOK PUBLISHING PRODUCTION TRACKER")
    title_cell.fill = PatternFill("solid", fgColor=navy)
    title_cell.font = Font(name="Aptos Display", size=16, bold=True, color=white)
    title_cell.alignment = Alignment(vertical="center")
    ws.row_dimensions[1].height = 32

    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=len(columns))
    subtitle = ws.cell(2, 1, f"Generated {today.isoformat()}  |  {len(tasks)} task(s)  |  Color-coded deadlines and statuses")
    subtitle.fill = PatternFill("solid", fgColor=light_blue)
    subtitle.font = Font(italic=True, color=navy, size=10)
    subtitle.alignment = Alignment(vertical="center")
    ws.row_dimensions[2].height = 22

    header_row = 4
    for col_num, (label, _) in enumerate(columns, start=1):
        cell = ws.cell(header_row, col_num, label)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = thin_border
    ws.row_dimensions[header_row].height = 28

    for row_num, task in enumerate(tasks, start=header_row + 1):
        for col_num, (_, key) in enumerate(columns, start=1):
            value = task.get(key, "")
            if key == "days_until_deadline" and value != "":
                try:
                    value = int(value)
                except ValueError:
                    pass
            cell = ws.cell(row_num, col_num, value)
            cell.alignment = Alignment(vertical="top", wrap_text=(key in {"task", "dependency", "notes", "reminder"}))
            cell.border = thin_border
            if (row_num - header_row) % 2 == 0:
                cell.fill = PatternFill("solid", fgColor="F8FAFD")

        # Color the reminder cell by urgency.
        reminder_cell = ws.cell(row_num, 8)
        reminder = str(task.get("reminder", "")).lower()
        if "overdue" in reminder or "due today" in reminder:
            reminder_cell.fill = PatternFill("solid", fgColor=red_fill)
            reminder_cell.font = Font(color=red_font, bold=True)
        elif "tomorrow" in reminder:
            reminder_cell.fill = PatternFill("solid", fgColor=orange_fill)
            reminder_cell.font = Font(color=orange_font, bold=True)
        elif "in 2 days" in reminder or "in 3 days" in reminder:
            reminder_cell.fill = PatternFill("solid", fgColor=yellow_fill)
            reminder_cell.font = Font(color=yellow_font, bold=True)
        elif "completed" in reminder:
            reminder_cell.fill = PatternFill("solid", fgColor=green_fill)
            reminder_cell.font = Font(color=green_font, bold=True)
        elif "missing/invalid" in reminder:
            reminder_cell.fill = PatternFill("solid", fgColor=purple_fill)
            reminder_cell.font = Font(color=purple_font, bold=True)

        # Color the status cell independently.
        status_cell = ws.cell(row_num, 10)
        status = str(task.get("status", "")).strip().lower()
        if status in {"blocked", "on hold"}:
            status_cell.fill = PatternFill("solid", fgColor=red_fill)
            status_cell.font = Font(color=red_font, bold=True)
        elif status in {"pending", "not started", "in review"}:
            status_cell.fill = PatternFill("solid", fgColor=yellow_fill)
            status_cell.font = Font(color=yellow_font, bold=True)
        elif status == "in progress":
            status_cell.fill = PatternFill("solid", fgColor=light_blue)
            status_cell.font = Font(color=blue, bold=True)
        elif status in {"completed", "approved", "published", "cancelled"}:
            status_cell.fill = PatternFill("solid", fgColor=green_fill)
            status_cell.font = Font(color=green_font, bold=True)

        # Color the priority cell.
        priority_cell = ws.cell(row_num, 9)
        priority = str(task.get("priority", "")).lower()
        if priority == "critical":
            priority_cell.fill = PatternFill("solid", fgColor=red_fill)
            priority_cell.font = Font(color=red_font, bold=True)
        elif priority == "high":
            priority_cell.fill = PatternFill("solid", fgColor=orange_fill)
            priority_cell.font = Font(color=orange_font, bold=True)
        elif priority == "medium":
            priority_cell.fill = PatternFill("solid", fgColor=yellow_fill)
            priority_cell.font = Font(color=yellow_font)
        elif priority == "low":
            priority_cell.fill = PatternFill("solid", fgColor=green_fill)
            priority_cell.font = Font(color=green_font)

    ws.freeze_panes = "A5"
    ws.auto_filter.ref = f"A4:{get_column_letter(len(columns))}{max(header_row + len(tasks), header_row)}"
    widths = [25, 14, 30, 21, 20, 16, 12, 27, 14, 17, 25, 34]
    for index, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(index)].width = width
    ws.sheet_view.showGridLines = False

    # Summary sheet
    summary = wb.create_sheet("Summary")
    summary.merge_cells("A1:B1")
    summary["A1"] = "PRODUCTION SUMMARY"
    summary["A1"].fill = PatternFill("solid", fgColor=navy)
    summary["A1"].font = Font(size=16, bold=True, color=white)
    summary["A1"].alignment = Alignment(vertical="center")
    summary.row_dimensions[1].height = 32

    active = [t for t in tasks if t["_completed"] != "yes"]
    completed = [t for t in tasks if t["_completed"] == "yes"]
    overdue = [t for t in active if t["days_until_deadline"] and int(t["days_until_deadline"]) < 0]
    due_today = [t for t in active if t["days_until_deadline"] == "0"]
    due_soon = [t for t in active if t["days_until_deadline"] in {"1", "2", "3"}]
    missing_dates = [t for t in active if not t["days_until_deadline"]]
    blocked = [t for t in active if t["status"].lower() in {"blocked", "on hold"}]
    unique_books = {t["book_title"].casefold() for t in tasks}
    metrics = [
        ("Report date", today.isoformat()),
        ("Books represented", len(unique_books)),
        ("Total tasks retained", len(tasks)),
        ("Active tasks", len(active)),
        ("Completed/cancelled tasks", len(completed)),
        ("Overdue tasks", len(overdue)),
        ("Due today", len(due_today)),
        ("Due within 3 days", len(due_soon)),
        ("Missing/invalid deadlines", len(missing_dates)),
        ("Blocked/on-hold tasks", len(blocked)),
        ("Data warnings", len(warnings)),
    ]
    for r, (label, value) in enumerate(metrics, start=3):
        a = summary.cell(r, 1, label)
        b = summary.cell(r, 2, value)
        a.font = Font(bold=True, color=navy)
        a.fill = PatternFill("solid", fgColor=light_grey)
        b.fill = PatternFill("solid", fgColor="FFFFFF")
        a.border = thin_border
        b.border = thin_border
        if label in {"Overdue tasks", "Due today"} and int(value) > 0:
            b.fill = PatternFill("solid", fgColor=red_fill)
            b.font = Font(color=red_font, bold=True)
        elif label == "Due within 3 days" and int(value) > 0:
            b.fill = PatternFill("solid", fgColor=yellow_fill)
            b.font = Font(color=yellow_font, bold=True)

    summary["A16"] = "Data quality notes"
    summary["A16"].fill = header_fill
    summary["A16"].font = header_font
    summary.merge_cells("A16:B16")
    if warnings:
        for index, warning in enumerate(warnings, start=17):
            summary.cell(index, 1, warning)
            summary.merge_cells(start_row=index, start_column=1, end_row=index, end_column=2)
            summary.cell(index, 1).alignment = Alignment(wrap_text=True, vertical="top")
    else:
        summary["A17"] = "No unparseable dates found."
        summary.merge_cells("A17:B17")
    summary.column_dimensions["A"].width = 32
    summary.column_dimensions["B"].width = 24
    summary.sheet_view.showGridLines = False
    summary.freeze_panes = "A3"

    wb.save(output_path)
    return True


def main() -> int:
    parser = argparse.ArgumentParser(
        description="Clean and prioritize a book-publishing production task CSV."
    )
    parser.add_argument("csv_path", nargs="?", default="input_tasks.csv",
                        help="Input CSV path (default: input_tasks.csv)")
    parser.add_argument("--output-dir", default="publishing_reports",
                        help="Folder for generated reports (default: publishing_reports)")
    parser.add_argument("--today", help="Override today's date for testing (YYYY-MM-DD)")
    args = parser.parse_args()

    if args.today:
        try:
            today = date.fromisoformat(args.today)
        except ValueError:
            print("Error: --today must use YYYY-MM-DD format.", file=sys.stderr)
            return 2
    else:
        today = date.today()

    csv_path = Path(args.csv_path)
    output_dir = Path(args.output_dir)

    try:
        tasks, warnings = read_tasks(csv_path, today)
        sorted_tasks = sort_tasks(tasks)
        output_dir.mkdir(parents=True, exist_ok=True)

        excel_path = output_dir / "publishing_tasks_report.xlsx"
        summary_path = output_dir / "publishing_summary.txt"
        summary = write_summary(sorted_tasks, warnings, summary_path, today)
        excel_created = write_excel_report(sorted_tasks, warnings, excel_path, today)

        print(summary)
        print("\nFiles created:")
        if excel_created:
            print(f"- Formatted Excel report: {excel_path.resolve()}")
        else:
            print("- Formatted Excel report not created; install openpyxl with: python -m pip install openpyxl")
        print(f"- Summary: {summary_path.resolve()}")
        print("\nNote: reminders are calculated when this script runs; this version does not send notifications.")
        return 0
    except (OSError, ValueError, csv.Error) as exc:
        print(f"Error: {exc}", file=sys.stderr)
        return 1


if __name__ == "__main__":
    raise SystemExit(main())
